from __future__ import annotations

import io
import math
import os
import shutil
import subprocess
import tempfile
import wave
import zipfile
from pathlib import Path

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'whut_prime.settings')

import django

django.setup()

from django.core.files.base import ContentFile
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from portal.models import RecruitmentApplication, RecruitmentAttachment


def make_image(path: Path) -> None:
    image = Image.new('RGB', (1200, 800), '#eaf5f6')
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1200, 116), fill='#0e3e56')
    draw.rectangle((70, 190, 1130, 690), outline='#168bb5', width=6)
    draw.ellipse((145, 270, 405, 530), fill='#2de2a6', outline='#0e3e56', width=8)
    draw.rectangle((500, 300, 1030, 350), fill='#168bb5')
    draw.rectangle((500, 390, 900, 440), fill='#7bb9c9')
    draw.rectangle((500, 480, 780, 530), fill='#b7dce2')
    draw.text((72, 38), 'PRIME MATERIAL PREVIEW', fill='#ffffff')
    draw.text((178, 360), 'IMG', fill='#0e3e56')
    image.save(path, 'PNG')


def make_pdf(path: Path) -> None:
    document = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    for page in range(1, 4):
        document.setFillColorRGB(0.055, 0.243, 0.337)
        document.setFont('Helvetica-Bold', 24)
        document.drawString(54, height - 70, f'PRIME PDF preview page {page}')
        document.setStrokeColorRGB(0.086, 0.545, 0.71)
        document.setLineWidth(2)
        document.line(54, height - 92, width - 54, height - 92)
        document.setFillColorRGB(0.18, 0.28, 0.33)
        document.setFont('Helvetica', 13)
        for line in range(7):
            document.drawString(70, height - 155 - line * 34, f'Server-side PDF rendering test · line {line + 1}')
        document.setFillColorRGB(0.18, 0.75, 0.61)
        document.circle(width / 2, height / 2, 70, fill=1, stroke=0)
        document.setFillColorRGB(1, 1, 1)
        document.setFont('Helvetica-Bold', 20)
        document.drawCentredString(width / 2, height / 2 - 7, str(page))
        document.showPage()
    document.save()


def make_docx(path: Path) -> None:
    document = Document()
    for page in range(1, 3):
        document.add_heading(f'PRIME 文档预览测试 · 第 {page} 页', level=1)
        document.add_paragraph('这是一份模拟报名材料，用来检查服务器端 Word 转 PDF 和多页纵向阅读。')
        document.add_paragraph('包含段落、表格和分页，点击左侧缩略图可以跳转到对应页面。')
        table = document.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        for cell, value in zip(table.rows[0].cells, ('项目', '状态', '说明')):
            cell.text = value
        for row in range(1, 5):
            cells = table.add_row().cells
            cells[0].text = f'测试项目 {row}'
            cells[1].text = '已准备'
            cells[2].text = '用于预览验证'
        if page == 1:
            document.add_page_break()
    document.save(path)


def make_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '报名材料'
    sheet.append(['编号', '测试项目', '结果'])
    for row in range(1, 9):
        sheet.append([row, f'格式预览项目 {row}', '待检查'])
    detail = workbook.create_sheet('说明')
    detail.append(['字段', '内容'])
    detail.append(['用途', '检查表格转 PDF 后的排版'])
    detail.append(['页面', '左侧缩略图 + 右侧纵向阅读'])
    workbook.save(path)


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    for index, title in enumerate(('PPT 预览测试', '服务器转换', '缩略图与纵向阅读'), start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(238, 248, 249)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, 850000)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(14, 62, 86)
        box.line.fill.background()
        text_box = slide.shapes.add_textbox(900000, 1300000, 8500000, 900000)
        text_box.text_frame.text = title
        text_box.text_frame.paragraphs[0].font.size = Pt(30)
        body = slide.shapes.add_textbox(900000, 2600000, 8500000, 1700000)
        body.text_frame.text = f'第 {index} 页\n由服务器端 PowerPoint 转成 PDF 进行预览'
        body.text_frame.paragraphs[0].font.size = Pt(22)
    presentation.save(path)


def make_wav(path: Path) -> None:
    sample_rate = 44100
    seconds = 4
    with wave.open(str(path), 'wb') as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(sample_rate)
        frames = bytearray()
        for index in range(sample_rate * seconds):
            value = int(0.22 * 32767 * math.sin(2 * math.pi * 440 * index / sample_rate))
            frames.extend(value.to_bytes(2, byteorder='little', signed=True))
        audio.writeframes(frames)


def make_mp4(path: Path, image: Path) -> None:
    ffmpeg = shutil.which('ffmpeg')
    if not ffmpeg:
        raise RuntimeError('ffmpeg 未安装，无法生成测试视频。')
    subprocess.run(
        [ffmpeg, '-y', '-loop', '1', '-i', str(image), '-t', '4', '-r', '24', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', str(path)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def main() -> None:
    with tempfile.TemporaryDirectory(prefix='prime-preview-demo-') as directory:
        root = Path(directory)
        files = {
            'preview-demo.png': root / 'preview-demo.png',
            'preview-demo.pdf': root / 'preview-demo.pdf',
            'preview-demo.docx': root / 'preview-demo.docx',
            'preview-demo.xlsx': root / 'preview-demo.xlsx',
            'preview-demo.pptx': root / 'preview-demo.pptx',
            'preview-demo.wav': root / 'preview-demo.wav',
            'preview-demo.mp4': root / 'preview-demo.mp4',
            'preview-demo.txt': root / 'preview-demo.txt',
        }
        make_image(files['preview-demo.png'])
        make_pdf(files['preview-demo.pdf'])
        make_docx(files['preview-demo.docx'])
        make_xlsx(files['preview-demo.xlsx'])
        make_pptx(files['preview-demo.pptx'])
        make_wav(files['preview-demo.wav'])
        make_mp4(files['preview-demo.mp4'], files['preview-demo.png'])
        files['preview-demo.txt'].write_text('文本预览测试\n支持直接阅读纯文本文件。\n', encoding='utf-8')

        zip_path = root / 'preview-demo.zip'
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, path in files.items():
                archive.write(path, f'报名材料/{name}')
            archive.writestr('报名材料/目录/README.txt', 'ZIP 子目录预览测试\n')
        files['preview-demo.zip'] = zip_path

        email = 'preview-demo-20260918@example.com'
        application, created = RecruitmentApplication.objects.get_or_create(
            email=email,
            defaults={
                'name': '格式预览测试样例',
                'gender': 'male',
                'qq': '2609180000',
                'wechat': 'prime-preview-demo',
                'phone': '13800000000',
                'college': '人工智能学院',
                'major_class': '自动化类2601班',
                'intended_groups': ['algorithm', 'operations'],
                'primary_choice': 'algorithm',
                'accepts_adjustment': True,
                'second_choice': 'operations',
                'introduction': '用于验证多种报名材料的在线预览效果。',
                'honors': '校级科技创新项目一等奖',
                'roles': '班级学习委员',
                'technical_foundation': 'Python、C++、基础电路',
                'experience': '完成过机器人视觉和流水灯相关项目。',
                'availability': '',
                'consent': True,
            },
        )
        if not created:
            application.attachments.all().delete()

        for original_name, path in files.items():
            attachment = RecruitmentAttachment(application=application, original_name=original_name)
            attachment.file.save(original_name, ContentFile(path.read_bytes()), save=True)

        print(f'application_id={application.pk}')
        print(f'application_no={application.application_no}')
        print(f'attachments={application.attachments.count()}')


if __name__ == '__main__':
    main()
