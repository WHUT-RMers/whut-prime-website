from io import BytesIO
import zipfile

from openpyxl import load_workbook

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core import mail
from django.test import TestCase
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone

from .models import (NewsArticle, NewsCoverSlide, NewsImage, RecruitmentApplication,
                     RecruitmentEmailVerification, RecruitmentSettings)


@override_settings(RECRUITMENT_EMAIL_VERIFICATION_ENABLED=False)
class PortalApiTests(TestCase):
    def setUp(self):
        self.article = NewsArticle.objects.create(
            title='测试资讯', summary='测试摘要', category='赛报', body='<p>正文</p>',
            status=NewsArticle.Status.PUBLISHED, published_at=timezone.now(), cover='news/test.jpg',
            cover_caption='机器人在赛场进行调试',
        )

    def photo_upload(self, filename='portrait.jpg'):
        photo_buffer = BytesIO()
        Image.new('RGB', (480, 640), '#dbeef2').save(photo_buffer, 'JPEG')
        return SimpleUploadedFile(filename, photo_buffer.getvalue(), content_type='image/jpeg')

    def test_public_news_and_deduplicated_view(self):
        item = self.client.get(reverse('portal:news-list')).json()['items'][0]
        self.assertEqual(item['slug'], self.article.slug)
        self.assertEqual(item['cover_images'][0]['caption'], '机器人在赛场进行调试')
        url = reverse('portal:news-view', args=[self.article.slug])
        self.client.post(url)
        self.client.post(url)
        self.article.refresh_from_db()
        self.assertEqual(self.article.view_count, 1)

    def test_closed_recruitment_rejects_submission(self):
        RecruitmentSettings.current()
        response = self.client.post(reverse('portal:recruitment-submit'), {})
        self.assertEqual(response.status_code, 403)

    def test_open_recruitment_accepts_pdf(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '张同学', 'gender': 'male', 'qq': '12345678', 'wechat': 'prime-zhang', 'email': 'zhang@example.com', 'phone': '13800000000',
            'college': '计算机学院', 'major_class': '软件工程 2301 班',
            'intended_groups': ['algorithm'], 'introduction': '热爱机器人', 'honors': '无', 'roles': '无', 'technical_foundation': 'Python 基础', 'experience': '无', 'availability': '每周 12 小时', 'consent': 'true',
            'photo': self.photo_upload('zhang.jpg'),
            'resume': SimpleUploadedFile('resume.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
        })
        self.assertEqual(response.status_code, 201, response.content.decode())

    def test_open_recruitment_accepts_without_attachment(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '无附件同学', 'gender': 'female', 'qq': '12345679', 'wechat': 'prime-no-file', 'email': 'no-file@example.com', 'phone': '13800000009',
            'college': '计算机学院', 'major_class': '软件工程 2302 班', 'primary_choice': 'algorithm',
            'introduction': '暂时没有附件', 'honors': '无', 'roles': '无', 'technical_foundation': '正在学习 Python',
            'experience': '无', 'availability': '每周 8 小时', 'consent': 'true',
            'photo': self.photo_upload('no-file.jpg'),
        })
        self.assertEqual(response.status_code, 201, response.content.decode())
        application = RecruitmentApplication.objects.get(name='无附件同学')
        self.assertFalse(application.resume)
        self.assertEqual(application.attachments.count(), 0)

    def test_latest_submission_replaces_previous_application_for_email(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        first = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '第一次提交', 'gender': 'male', 'qq': '10001', 'wechat': 'latest-demo', 'email': 'latest@example.com', 'phone': '13800000001',
            'college': '计算机学院', 'major_class': '软件工程 2301 班', 'primary_choice': 'algorithm',
            'introduction': '第一次内容', 'honors': '无', 'roles': '无', 'technical_foundation': 'Python', 'experience': '第一次经历', 'consent': 'true',
            'attachments': SimpleUploadedFile('first.txt', b'first'),
        })
        self.assertEqual(first.status_code, 201, first.content.decode())
        original = RecruitmentApplication.objects.get(email='latest@example.com')
        original_id = original.pk

        second = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '最后一次提交', 'gender': 'female', 'qq': '10002', 'wechat': 'latest-demo', 'email': 'latest@example.com', 'phone': '13800000002',
            'college': '自动化学院', 'major_class': '自动化 2301 班', 'primary_choice': 'electrical',
            'introduction': '最后一次内容', 'honors': '一等奖', 'roles': '无', 'technical_foundation': 'C++', 'experience': '最后一次经历', 'consent': 'true',
            'attachments': SimpleUploadedFile('latest.txt', b'latest'),
        })
        self.assertEqual(second.status_code, 200, second.content.decode())
        self.assertEqual(RecruitmentApplication.objects.filter(email='latest@example.com').count(), 1)
        application = RecruitmentApplication.objects.get(email='latest@example.com')
        self.assertEqual(application.pk, original_id)
        self.assertEqual(application.name, '最后一次提交')
        self.assertEqual(application.primary_choice, 'electrical')
        self.assertEqual(application.attachments.count(), 1)
        self.assertEqual(application.attachments.get().original_name, 'latest.txt')

    def test_open_recruitment_accepts_optional_photo(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        photo_buffer = BytesIO()
        Image.new('RGB', (480, 640), '#dbeef2').save(photo_buffer, 'JPEG')
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '照片同学', 'gender': 'male', 'qq': '12345681', 'wechat': 'prime-photo', 'email': 'photo@example.com', 'phone': '13800000011',
            'college': '计算机学院', 'major_class': '软件工程 2304 班', 'primary_choice': 'algorithm',
            'introduction': '热爱机器人', 'honors': '无', 'roles': '无', 'technical_foundation': 'Python 基础', 'experience': '无',
            'availability': '每周 8 小时', 'consent': 'true',
            'photo': SimpleUploadedFile('portrait.jpg', photo_buffer.getvalue(), content_type='image/jpeg'),
        })
        self.assertEqual(response.status_code, 201, response.content.decode())
        self.assertTrue(RecruitmentApplication.objects.get(name='照片同学').photo)

    def test_open_recruitment_accepts_without_photo(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '缺照片同学', 'gender': 'male', 'qq': '12345682', 'wechat': 'prime-no-photo', 'email': 'no-photo@example.com', 'phone': '13800000012',
            'college': '计算机学院', 'major_class': '软件工程 2305 班', 'primary_choice': 'algorithm',
            'introduction': '热爱机器人', 'honors': '无', 'roles': '无', 'technical_foundation': 'Python 基础', 'experience': '无',
            'availability': '每周 8 小时', 'consent': 'true',
        })
        self.assertEqual(response.status_code, 201, response.content.decode())
        self.assertFalse(RecruitmentApplication.objects.get(name='缺照片同学').photo)

    def test_open_recruitment_requires_all_profile_text(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '缺少经历同学', 'gender': 'female', 'qq': '12345680', 'wechat': 'prime-missing', 'email': 'missing@example.com', 'phone': '13800000010',
            'college': '计算机学院', 'major_class': '软件工程 2303 班', 'primary_choice': 'algorithm',
            'introduction': '热爱机器人', 'honors': '无', 'roles': '无', 'technical_foundation': 'Python 基础',
            'experience': '', 'availability': '每周 8 小时', 'consent': 'true',
            'photo': self.photo_upload('missing.jpg'),
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn('experience', response.json()['errors'])

    def test_open_recruitment_accepts_any_supported_attachment(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '材料同学', 'gender': 'male', 'qq': '12345678', 'wechat': 'prime-file', 'email': 'file@example.com', 'phone': '13800000002',
            'college': '计算机学院', 'major_class': '软件工程 2301 班', 'primary_choice': 'algorithm',
            'introduction': '用 Word 材料报名', 'honors': '无', 'roles': '无', 'technical_foundation': 'Office 基础', 'experience': '无', 'availability': '每周 12 小时', 'consent': 'true',
            'photo': self.photo_upload('file.jpg'),
            'attachments': SimpleUploadedFile('portfolio.docx', b'not-a-real-docx-but-a-valid-upload', content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
        })
        self.assertEqual(response.status_code, 201, response.content.decode())
        self.assertEqual(RecruitmentApplication.objects.get(name='材料同学').attachments.count(), 1)

    def test_download_archives_all_uploaded_attachments_without_pdf(self):
        application = RecruitmentApplication.objects.create(
            name='下载材料同学', qq='12345678', wechat='download-file', email='download@example.com', phone='13800000004',
            college='计算机学院', major_class='软件工程 2301 班', primary_choice='algorithm', intended_groups=['algorithm'],
            introduction='下载非 PDF 材料', availability='每周 12 小时', consent=True,
        )
        application.attachments.create(
            original_name='portfolio.docx', file=SimpleUploadedFile('portfolio.docx', b'document-content'),
        )
        application.attachments.create(
            original_name='note.md', file=SimpleUploadedFile('note.md', b'# note'),
        )
        user = get_user_model().objects.create_superuser('archive-admin', 'archive-admin@example.com', 'safe-password-123')
        self.client.force_login(user)
        response = self.client.get(reverse('portal:resume-download', args=[application.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['Content-Type'], 'application/zip')
        with zipfile.ZipFile(BytesIO(b''.join(response.streaming_content))) as archive:
            names = archive.namelist()
        self.assertTrue(any(name.endswith('/portfolio.docx') for name in names))
        self.assertTrue(any(name.endswith('/note.md') for name in names))

    @override_settings(
        RECRUITMENT_EMAIL_VERIFICATION_ENABLED=True,
        EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
        EMAIL_HOST_USER='test-sender@qq.com', EMAIL_HOST_PASSWORD='test-auth-code',
    )
    def test_recruitment_requires_a_verified_email_when_enabled(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        email = 'verified@example.com'
        send = self.client.post(reverse('portal:recruitment-send-email-code'), {'email': email})
        self.assertEqual(send.status_code, 200, send.content.decode())
        self.assertEqual(len(mail.outbox), 1)
        import re
        token = re.search(r'[?&]token=([^ &\n]+)', mail.outbox[0].body).group(1)
        verify = self.client.get(reverse('portal:recruitment-email-verify'), {'token': token})
        self.assertEqual(verify.status_code, 200, verify.content.decode())
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '验证同学', 'gender': 'female', 'qq': '12345678', 'wechat': 'verified-prime', 'email': email, 'phone': '13800000000',
            'college': '计算机学院', 'major_class': '软件工程 2301 班', 'primary_choice': 'algorithm',
            'introduction': '热爱机器人', 'honors': '无', 'roles': '无', 'technical_foundation': 'C++ 基础', 'experience': '无', 'availability': '每周 12 小时', 'consent': 'true',
            'photo': self.photo_upload('verified.jpg'),
            'attachments': SimpleUploadedFile('resume.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
        })
        self.assertEqual(response.status_code, 201, response.content.decode())
        self.assertTrue(RecruitmentApplication.objects.get(name='验证同学').email_verified)

    @override_settings(
        RECRUITMENT_EMAIL_VERIFICATION_ENABLED=True,
        EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
        EMAIL_HOST_USER='test-sender@qq.com', EMAIL_HOST_PASSWORD='test-auth-code',
    )
    def test_verified_application_accepts_multiple_non_pdf_attachments(self):
        import re
        email = 'materials-verified@example.com'
        self.client.post(reverse('portal:recruitment-send-email-code'), {'email': email})
        token = re.search(r'[?&]token=([^ &\n]+)', mail.outbox[0].body).group(1)
        self.assertEqual(self.client.get(reverse('portal:recruitment-email-verify'), {'token': token}).status_code, 200)
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '多材料同学', 'gender': 'male', 'qq': '87654321', 'wechat': 'multi-material', 'email': email, 'phone': '13800000003',
            'college': '自动化学院', 'major_class': '自动化 2301 班', 'primary_choice': 'algorithm',
            'introduction': '使用多种材料报名', 'honors': '无', 'roles': '无', 'technical_foundation': 'Python 基础', 'experience': '无', 'availability': '每周 12 小时', 'consent': 'true',
            'photo': self.photo_upload('multi-material.jpg'),
            'attachments': [
                SimpleUploadedFile('portfolio.docx', b'document'),
                SimpleUploadedFile('note.md', b'# note'),
                SimpleUploadedFile('project.zip', b'archive'),
            ],
        })
        self.assertEqual(response.status_code, 201, response.content.decode())
        self.assertEqual(RecruitmentApplication.objects.get(name='多材料同学').attachments.count(), 3)

    @override_settings(
        RECRUITMENT_EMAIL_VERIFICATION_ENABLED=True,
        EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
        EMAIL_HOST_USER='test-sender@qq.com', EMAIL_HOST_PASSWORD='test-auth-code',
    )
    def test_sixth_daily_email_code_request_is_rejected_before_sending(self):
        from datetime import timedelta
        from django.contrib.auth.hashers import make_password
        now = timezone.now()
        RecruitmentEmailVerification.objects.create(
            email='limited@example.com', code_hash=make_password('123456'),
            expires_at=now + timedelta(minutes=10), last_sent_at=now - timedelta(minutes=2),
            sent_count=5, sent_on=timezone.localdate(),
        )
        response = self.client.post(reverse('portal:recruitment-send-email-code'), {'email': 'limited@example.com'})
        self.assertEqual(response.status_code, 429, response.content.decode())
        self.assertEqual(len(mail.outbox), 0)

    @override_settings(
        RECRUITMENT_EMAIL_VERIFICATION_ENABLED=True,
        EMAIL_HOST_USER='test-sender@qq.com', EMAIL_HOST_PASSWORD='test-auth-code',
    )
    def test_verified_applicant_can_modify_pending_application_once(self):
        application = RecruitmentApplication.objects.create(
            name='可修改同学', qq='111', wechat='editable', email='editable@example.com', phone='13800000001',
            gender='male', college='自动化学院', major_class='自动化 2301 班', primary_choice='electrical', intended_groups=['electrical'],
            introduction='原始介绍', availability='每周 10 小时', consent=True,
            photo=self.photo_upload('editable.jpg'),
            resume=SimpleUploadedFile('resume.pdf', b'%PDF-1.4 editable', content_type='application/pdf'),
        )
        # 链接点击后：数据库记录 verified_at 已写入（新流程以记录为准，不再依赖会话）
        from datetime import timedelta
        from django.contrib.auth.hashers import make_password
        now = timezone.now()
        RecruitmentEmailVerification.objects.create(
            email=application.email, code_hash=make_password('token'),
            expires_at=now + timedelta(minutes=10), last_sent_at=now - timedelta(minutes=5),
            verified_at=now,
        )
        status = self.client.post(reverse('portal:recruitment-application-status'), {'email': application.email})
        self.assertTrue(status.json()['application']['can_edit'])
        response = self.client.post(reverse('portal:recruitment-update', args=[application.pk]), {
            'name': application.name, 'gender': application.gender, 'qq': application.qq, 'wechat': application.wechat, 'email': application.email, 'phone': application.phone,
            'college': application.college, 'major_class': application.major_class, 'primary_choice': 'algorithm',
            'introduction': '更新后的介绍', 'honors': '无', 'roles': '无', 'technical_foundation': 'Python 基础', 'experience': '更新后的项目经历', 'availability': application.availability, 'consent': 'true',
        })
        self.assertEqual(response.status_code, 200, response.content.decode())
        application.refresh_from_db()
        self.assertEqual(application.primary_choice, 'algorithm')
        self.assertEqual(application.modification_count, 1)
        self.assertEqual(application.applicant_revisions.count(), 0)

    def test_recruitment_choices_attachments_and_grouped_export(self):
        settings = RecruitmentSettings.current(); settings.is_open = True; settings.save()
        response = self.client.post(reverse('portal:recruitment-submit'), {
            'name': '李同学', 'gender': 'female', 'qq': '12345678', 'wechat': 'prime-li', 'email': 'li@example.com', 'phone': '13900000000',
            'college': '自动化学院', 'major_class': '自动化 2301 班', 'primary_choice': 'algorithm',
            'accepts_adjustment': 'true', 'second_choice': 'electrical', 'introduction': '热爱机器人', 'honors': '无', 'roles': '无', 'technical_foundation': 'C++ 基础', 'experience': '无',
            'availability': '每周 12 小时', 'consent': 'true',
            'photo': self.photo_upload('li.jpg'),
            'attachments': [SimpleUploadedFile('resume.pdf', b'%PDF-1.4 test', content_type='application/pdf'), SimpleUploadedFile('award.pdf', b'%PDF-1.4 award', content_type='application/pdf')],
        })
        self.assertEqual(response.status_code, 201, response.content.decode())
        application = RecruitmentApplication.objects.get(name='李同学')
        self.assertEqual(application.primary_choice, 'algorithm')
        self.assertEqual(application.second_choice, 'electrical')
        self.assertEqual(application.attachments.count(), 2)
        user = get_user_model().objects.create_superuser('recruit-exporter', 'recruit@example.com', 'safe-password-123')
        self.client.force_login(user)
        export = self.client.get(reverse('admin:portal_recruitment_export_all'))
        self.assertEqual(export.status_code, 200)
        with zipfile.ZipFile(BytesIO(b''.join(export.streaming_content))) as archive:
            names = archive.namelist()
            workbook = load_workbook(BytesIO(archive.read('报名信息.xlsx')))
        self.assertIn('报名信息.xlsx', names)
        self.assertTrue(any(name.startswith('算法组/') and name.endswith('resume.pdf') for name in names))
        self.assertEqual(workbook.sheetnames, ['报名汇总'])
        self.assertTrue(any('算法组' in str(row[0].value) for row in workbook['报名汇总'].iter_rows()))

    def test_staff_can_preview_recruitment_materials(self):
        application = RecruitmentApplication.objects.create(
            name='预览同学', qq='7654321', wechat='preview-prime', email='preview@example.com', phone='13700000000',
            college='计算机学院', major_class='软件工程 2301 班', primary_choice='algorithm', intended_groups=['algorithm'],
            introduction='用于验证后台附件预览。', availability='每周 10 小时', consent=True,
            resume=SimpleUploadedFile('resume.pdf', b'%PDF-1.4 preview', content_type='application/pdf'),
        )
        attachment = application.attachments.create(
            original_name='award.pdf',
            file=SimpleUploadedFile('award.pdf', b'%PDF-1.4 award', content_type='application/pdf'),
        )
        user = get_user_model().objects.create_superuser('preview-admin', 'preview-admin@example.com', 'safe-password-123')
        self.client.force_login(user)
        preview = self.client.get(reverse('admin:portal_recruitment_attachment_file', args=[application.pk, attachment.pk]))
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.headers['Content-Type'], 'application/pdf')
        self.assertNotIn('attachment', preview.headers.get('Content-Disposition', ''))
        self.assertEqual(
            self.client.get(reverse('admin:portal_recruitment_resume_preview', args=[application.pk])).status_code,
            200,
        )

    def test_staff_can_preview_application_form_and_navigate_filtered(self):
        photo_buffer = BytesIO()
        Image.new('RGB', (480, 640), '#dbeef2').save(photo_buffer, 'JPEG')
        first = RecruitmentApplication.objects.create(
            name='表单预览一号', qq='7654301', wechat='preview-one', email='preview-one@example.com', phone='13700000011',
            college='计算机学院', major_class='软件工程 2301 班', primary_choice='algorithm', intended_groups=['algorithm'],
            introduction='第一份报名表', availability='每周 10 小时', consent=True,
            photo=SimpleUploadedFile('first.jpg', photo_buffer.getvalue(), content_type='image/jpeg'),
        )
        second = RecruitmentApplication.objects.create(
            name='表单预览二号', qq='7654302', wechat='preview-two', email='preview-two@example.com', phone='13700000012',
            college='自动化学院', major_class='自动化 2302 班', primary_choice='algorithm', intended_groups=['algorithm'],
            introduction='第二份报名表', availability='每周 8 小时', consent=True,
        )
        user = get_user_model().objects.create_superuser('form-preview-admin', 'form-preview@example.com', 'safe-password-123')
        self.client.force_login(user)
        preview_url = reverse('admin:portal_recruitmentapplication_preview', args=[first.pk]) + '?status=pending'
        response = self.client.get(preview_url)
        self.assertEqual(response.status_code, 200, response.content.decode())
        self.assertContains(response, '报名表预览')
        self.assertContains(response, first.name)
        self.assertContains(response, '下一份')
        self.assertContains(response, reverse('admin:portal_recruitmentapplication_preview', args=[second.pk]))
        photo_response = self.client.get(reverse('admin:portal_recruitment_photo_preview', args=[first.pk]))
        self.assertEqual(photo_response.status_code, 200)
        self.assertEqual(photo_response.headers['Content-Type'], 'image/jpeg')

    def test_staff_can_preview_common_office_and_archive_materials(self):
        from docx import Document
        from openpyxl import Workbook

        application = RecruitmentApplication.objects.create(
            name='材料同学', qq='6654321', wechat='materials-prime', email='materials@example.com', phone='13600000000',
            college='自动化学院', major_class='自动化 2301 班', primary_choice='electrical', intended_groups=['electrical'],
            introduction='用于验证常见文件预览。', availability='每周 10 小时', consent=True,
            resume=SimpleUploadedFile('resume.pdf', b'%PDF-1.4 preview', content_type='application/pdf'),
        )
        document_buffer = BytesIO(); document = Document(); document.add_paragraph('Word 文档预览内容'); document.save(document_buffer)
        spreadsheet_buffer = BytesIO(); workbook = Workbook(); workbook.active.title = '报名信息'; workbook.active.append(['姓名', '组别']); workbook.active.append(['小王', '电控组']); workbook.save(spreadsheet_buffer)
        from PIL import Image
        image_buffer = BytesIO(); Image.new('RGB', (120, 80), '#2de2a6').save(image_buffer, 'JPEG')
        archive_buffer = BytesIO()
        with zipfile.ZipFile(archive_buffer, 'w') as archive:
            archive.writestr('作品/说明.txt', '作品说明')
            archive.writestr('作品/图片.jpg', image_buffer.getvalue())
            archive.writestr('作品/简历.docx', document_buffer.getvalue())
        attachments = [
            application.attachments.create(original_name='portfolio.docx', file=SimpleUploadedFile('portfolio.docx', document_buffer.getvalue())),
            application.attachments.create(original_name='works.xlsx', file=SimpleUploadedFile('works.xlsx', spreadsheet_buffer.getvalue())),
            application.attachments.create(original_name='materials.zip', file=SimpleUploadedFile('materials.zip', archive_buffer.getvalue())),
        ]
        image_attachment = application.attachments.create(
            original_name='poster.jpg', file=SimpleUploadedFile('poster.jpg', image_buffer.getvalue(), content_type='image/jpeg'),
        )
        user = get_user_model().objects.create_superuser('materials-admin', 'materials-admin@example.com', 'safe-password-123')
        self.client.force_login(user)
        expected_kinds = ['document', 'spreadsheet', 'archive']
        for attachment, expected_kind in zip(attachments, expected_kinds):
            response = self.client.get(reverse('admin:portal_recruitment_attachment_preview_data', args=[application.pk, attachment.pk]))
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()['kind'], expected_kind)
            if expected_kind == 'document':
                self.assertIn('Word 文档预览内容', response.json()['paragraphs'])
                self.assertNotIn('ocr', response.json())
            if expected_kind == 'archive':
                archive_preview = response.json()
                self.assertEqual(archive_preview['files'][0]['path'], '作品/说明.txt')
                entry = archive_preview['files'][0]
                entry_response = self.client.get(entry['preview'])
                self.assertEqual(entry_response.status_code, 200)
                self.assertEqual(entry_response['Content-Type'], 'text/plain')
                self.assertEqual(entry_response.content.decode(), '作品说明')
                image_entry = next(item for item in archive_preview['files'] if item['name'] == '图片.jpg')
                self.assertEqual(self.client.get(image_entry['preview']).headers['Content-Type'], 'image/jpeg')
                docx_entry = next(item for item in archive_preview['files'] if item['name'] == '简历.docx')
                nested_docx = self.client.get(docx_entry['dataPreview'])
                self.assertEqual(nested_docx.status_code, 200)
                self.assertIn('Word 文档预览内容', nested_docx.json()['paragraphs'])
        image_response = self.client.get(reverse('admin:portal_recruitment_attachment_preview_data', args=[application.pk, image_attachment.pk]))
        self.assertEqual(image_response.status_code, 200)
        self.assertEqual(image_response.json()['kind'], 'image')
        self.assertNotIn('ocr', image_response.json())

    def test_staff_can_upload_an_inline_body_image(self):
        user = get_user_model().objects.create_superuser('editor', 'editor@example.com', 'safe-password-123')
        self.client.force_login(user)
        buffer = BytesIO()
        Image.new('RGB', (1200, 675), '#2de2a6').save(buffer, 'JPEG')
        upload = SimpleUploadedFile('inline.jpg', buffer.getvalue(), content_type='image/jpeg')
        response = self.client.post(reverse('admin:portal_news_upload_inline_image'), {
            'image': upload, 'article_id': self.article.pk,
        })
        self.assertEqual(response.status_code, 200, response.content.decode())
        self.assertEqual(NewsImage.objects.filter(article=self.article).count(), 1)
        self.assertTrue(response.json()['url'].startswith('/api/news/'))
        self.assertIn('/content-images/', response.json()['url'])
        image_response = self.client.get(response.json()['url'])
        self.assertEqual(image_response.status_code, 200)
        self.assertEqual(image_response.headers['Content-Type'], 'image/jpeg')
        self.client.logout()
        self.assertEqual(self.client.get(response.json()['url']).status_code, 200)

    def test_new_article_editor_can_upload_image_before_first_save(self):
        user = get_user_model().objects.create_superuser('new-body-editor', 'new-body@example.com', 'safe-password-123')
        self.client.force_login(user)
        buffer = BytesIO()
        Image.new('RGB', (1000, 700), '#4da3ff').save(buffer, 'JPEG')
        upload = SimpleUploadedFile('draft-inline.jpg', buffer.getvalue(), content_type='image/jpeg')
        response = self.client.post(reverse('admin:portal_news_upload_inline_image'), {'image': upload})
        self.assertEqual(response.status_code, 200, response.content.decode())
        image = NewsImage.objects.get(pk=response.json()['id'])
        self.assertIsNone(image.article_id)
        self.assertEqual(self.client.get(response.json()['url']).status_code, 200)
        self.client.logout()
        self.assertEqual(self.client.get(response.json()['url']).status_code, 404)

    def test_news_editor_has_the_unified_carousel_manager(self):
        user = get_user_model().objects.create_superuser('carousel-editor', 'carousel@example.com', 'safe-password-123')
        self.client.force_login(user)
        response = self.client.get(reverse('admin:portal_newsarticle_change', args=[self.article.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'portal-carousel-input')
        self.assertContains(response, 'portal/carousel_manager.css')
        self.assertContains(response, 'portal/editor/prime-editor.js')
        self.assertContains(response, '官网实时预览')

    @override_settings(XIUMI_APP_ID='', XIUMI_APP_SECRET='')
    def test_news_editor_exposes_xiumi_setup_entry(self):
        user = get_user_model().objects.create_superuser('xiumi-setup', 'xiumi-setup@example.com', 'safe-password-123')
        self.client.force_login(user)
        editor = self.client.get(reverse('admin:portal_newsarticle_change', args=[self.article.pk]))
        self.assertContains(editor, '秀米编辑')
        setup = self.client.get(reverse('admin:portal_news_xiumi', args=[self.article.pk]))
        self.assertEqual(setup.status_code, 200)
        self.assertContains(setup, '尚未配置秀米同步应用')
        self.assertContains(setup, 'XIUMI_APP_SECRET')

    @override_settings(XIUMI_APP_ID='demo-app-id', XIUMI_APP_SECRET='server-only-secret')
    def test_xiumi_entry_creates_signed_bind_redirect_without_exposing_secret(self):
        user = get_user_model().objects.create_superuser('xiumi-bound', 'xiumi-bound@example.com', 'safe-password-123')
        self.client.force_login(user)
        response = self.client.get(reverse('admin:portal_news_xiumi', args=[self.article.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertIn('https://xiumi.us/auth/partner/bind?', response['Location'])
        self.assertIn('appid=demo-app-id', response['Location'])
        self.assertIn('partner_user_id=admin%3A', response['Location'])
        self.assertNotIn('server-only-secret', response['Location'])

    def test_news_add_page_opens_without_existing_images(self):
        user = get_user_model().objects.create_superuser('add-editor', 'add@example.com', 'safe-password-123')
        self.client.force_login(user)
        response = self.client.get(reverse('admin:portal_newsarticle_add'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id_carousel_uploads')
        self.assertContains(response, 'multiple')
        self.assertContains(response, 'portal/carousel_manager.js')
        self.assertContains(response, reverse('admin:portal_news_upload_inline_image'))
        expected = f'<div class="prime-editor-workbench" data-upload-url="{reverse("admin:portal_news_upload_inline_image")}"'
        self.assertIn(expected.encode(), response.content)

    def test_editor_appends_multiple_carousel_images_in_upload_order(self):
        user = get_user_model().objects.create_superuser('batch-editor', 'batch@example.com', 'safe-password-123')
        self.client.force_login(user)

        def image_upload(name, color):
            buffer = BytesIO()
            Image.new('RGB', (1200, 675), color).save(buffer, 'JPEG')
            return SimpleUploadedFile(name, buffer.getvalue(), content_type='image/jpeg')

        response = self.client.post(reverse('admin:portal_newsarticle_change', args=[self.article.pk]), {
            'title': self.article.title, 'slug': self.article.slug, 'summary': self.article.summary,
            'category': self.article.category, 'body': self.article.body, 'status': 'published',
            'published_at': timezone.localtime(self.article.published_at).strftime('%Y-%m-%dT%H:%M'),
            'image_focus': 'center', 'external_url': '',
            'carousel_uploads': [image_upload('first.jpg', '#e2a62d'), image_upload('second.jpg', '#2da6e2')],
            'carousel_upload_captions': ['第一张追加图片说明', '第二张追加图片说明'],
            'carousel_saved_caption_cover': '更新后的首图说明',
            '_save': '保存',
        })
        self.assertEqual(response.status_code, 302, response.content.decode())
        slides = NewsCoverSlide.objects.filter(article=self.article).order_by('sort_order')
        self.assertEqual(slides.count(), 2)
        self.article.refresh_from_db()
        self.assertEqual(self.article.cover_caption, '更新后的首图说明')
        self.assertEqual(list(slides.values_list('caption', flat=True)), ['第一张追加图片说明', '第二张追加图片说明'])
        with Image.open(slides[0].image) as first, Image.open(slides[1].image) as second:
            self.assertGreater(first.getpixel((0, 0))[0], first.getpixel((0, 0))[2])
            self.assertGreater(second.getpixel((0, 0))[2], second.getpixel((0, 0))[0])

    def test_editor_can_create_article_with_multiple_carousel_images(self):
        user = get_user_model().objects.create_superuser('new-editor', 'new@example.com', 'safe-password-123')
        self.client.force_login(user)

        def image_upload(name, color):
            buffer = BytesIO()
            Image.new('RGB', (1200, 675), color).save(buffer, 'JPEG')
            return SimpleUploadedFile(name, buffer.getvalue(), content_type='image/jpeg')

        response = self.client.post(reverse('admin:portal_newsarticle_add'), {
            'title': '后台新增的多图资讯', 'slug': '', 'summary': '用于验证新增流程。',
            'category': '日常', 'body': '<p>正文</p>', 'status': 'draft',
            'published_at': '', 'image_focus': 'center', 'external_url': '',
            'carousel_uploads': [image_upload('one.jpg', '#e2a62d'), image_upload('two.jpg', '#2da6e2')],
            'carousel_upload_captions': ['首图说明', '第二张说明'],
            '_save': '保存',
        })
        self.assertEqual(response.status_code, 302, response.content.decode())
        article = NewsArticle.objects.get(title='后台新增的多图资讯')
        self.assertTrue(article.cover.name)
        self.assertEqual(article.cover_caption, '首图说明')
        slide = NewsCoverSlide.objects.get(article=article)
        self.assertEqual(slide.caption, '第二张说明')

    def test_first_save_attaches_temporary_body_images(self):
        user = get_user_model().objects.create_superuser('attach-editor', 'attach@example.com', 'safe-password-123')
        self.client.force_login(user)

        def image_upload(name, color, size=(1200, 675)):
            buffer = BytesIO()
            Image.new('RGB', size, color).save(buffer, 'JPEG')
            return SimpleUploadedFile(name, buffer.getvalue(), content_type='image/jpeg')

        upload_response = self.client.post(reverse('admin:portal_news_upload_inline_image'), {
            'image': image_upload('body.jpg', '#2de2a6', (1000, 700)),
        })
        self.assertEqual(upload_response.status_code, 200)
        image = NewsImage.objects.get(pk=upload_response.json()['id'])
        body = f'<p>开头</p><figure class="image"><img src="{upload_response.json()["url"]}"></figure>'
        response = self.client.post(reverse('admin:portal_newsarticle_add'), {
            'title': '带正文插图的资讯', 'slug': '', 'summary': '用于验证图片归属。',
            'category': '日常', 'body': body, 'status': 'draft', 'published_at': '',
            'external_url': '', 'carousel_uploads': [image_upload('cover.jpg', '#4da3ff')], '_save': '保存',
        })
        self.assertEqual(response.status_code, 302, response.content.decode())
        image.refresh_from_db()
        self.assertEqual(image.article.title, '带正文插图的资讯')
from io import BytesIO

from PIL import Image
from django.contrib.auth import get_user_model
